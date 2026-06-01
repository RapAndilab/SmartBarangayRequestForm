import os
import uuid
from docx import Document
from django.conf import settings

def generate_document_file(data: dict) -> str:
    """
    Generate a DOCX file based on document_type and data fields.
    Returns the relative download URL to the saved file.
    """

    document_type = data.get("document_type", "")
    username = data.get("username", "")

    # Select template
    if document_type == "clearance":
        template_name = "clearance_template.docx"
    elif document_type == "residency":
        # template_name = "residency_template.docx"
        template_name = "indigency_template.docx"
    elif document_type == "certificate":
        template_name = "certificate_template.docx"
    else:
        raise ValueError("Unsupported document type")

    template_path = os.path.join(settings.BASE_DIR, "templates", template_name)
    doc = Document(template_path)

    # Replace placeholders in the document paragraphs
    def replace_placeholder_in_paragraph(paragraph, data):
        full_text = "".join(run.text for run in paragraph.runs)
        replaced = False

        for key, value in data.items():
            placeholder = f"{{{{{key}}}}}"
            if placeholder in full_text:
                full_text = full_text.replace(placeholder, str(value))
                replaced = True

        if replaced:
            # Clear all runs but keep formatting in first run
            paragraph.runs[0].text = full_text
            for run in paragraph.runs[1:]:
                run.text = ""

    for para in doc.paragraphs:
        replace_placeholder_in_paragraph(para, data)

    # Generate unique filename
    unique_str = uuid.uuid4().hex
    output_filename = f"generated_{document_type}_{username}_{unique_str}.docx"
    output_path = os.path.join(settings.MEDIA_ROOT, output_filename)

    # Ensure the media directory exists (a fresh server may not have it yet).
    os.makedirs(settings.MEDIA_ROOT, exist_ok=True)

    # Save file
    doc.save(output_path)

    # Return download URL relative to MEDIA_URL
    return os.path.join(settings.MEDIA_URL, output_filename)
